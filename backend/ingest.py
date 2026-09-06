"""Document ingestion — add a client's real files to their database entry.

Accepts .docx, .pdf and .txt. Text is extracted on upload so an uploaded
document is swept exactly like the seeded ones: no separate path, no special
casing in the impact engine.
"""
import io
import json
import os
import re
from datetime import date

from pypdf import PdfReader
import docx as docxlib

DATA = os.path.join(os.path.dirname(__file__), "data")
CLIENTS = os.path.join(DATA, "clients", "clients.json")
UPLOADS = os.path.join(DATA, "uploads")

SUPPORTED = {".docx", ".pdf", ".txt"}


def extract_text(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
    elif ext == ".docx":
        d = docxlib.Document(io.BytesIO(content))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:                      # clauses often sit in tables
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    elif ext == ".txt":
        text = content.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {ext}. Use .docx, .pdf or .txt")

    for bad, good in (("‑", "-"), ("–", "-"), ("—", " - "),
                      ("‘", "'"), ("’", "'"),
                      ("“", '"'), ("”", '"'), (" ", " ")):
        text = text.replace(bad, good)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def load_clients() -> list:
    with open(CLIENTS, encoding="utf-8") as f:
        return json.load(f)


def save_clients(data: list) -> None:
    with open(CLIENTS, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def add_document(client_id: str, filename: str, content: bytes,
                 description: str = "") -> dict:
    """Extract, store and attach a document to a client."""
    text = extract_text(filename, content)
    if not text:
        raise ValueError("No text could be extracted from that file.")

    clients = load_clients()
    client = next((c for c in clients if c["id"] == client_id), None)
    if not client:
        raise ValueError(f"Unknown client: {client_id}")

    os.makedirs(UPLOADS, exist_ok=True)
    doc_id = f"{client_id}u{len([d for d in client['documents'] if 'u' in d['id']]) + 1}"
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", filename)
    with open(os.path.join(UPLOADS, f"{doc_id}_{safe}"), "wb") as f:
        f.write(content)

    doc = {
        "id": doc_id,
        "name": description.strip() or os.path.splitext(filename)[0],
        "added": date.today().isoformat(),
        "text": text,
        "uploaded": True,
        "source_file": filename,
    }
    client["documents"].append(doc)
    save_clients(clients)
    return {"client_id": client_id, "client_name": client["name"],
            "document": {k: doc[k] for k in ("id", "name", "added", "source_file")},
            "chars_extracted": len(text),
            "preview": text[:400]}


def add_client(name: str, practice: str, business: str = "",
               headcount: int = 0, notes: str = "") -> dict:
    clients = load_clients()
    cid = f"c{max([int(c['id'][1:]) for c in clients if c['id'][1:].isdigit()] + [0]) + 1}"
    client = {
        "id": cid, "name": name.strip(), "practice": practice.strip() or "General",
        "risk": "None", "updated": date.today().isoformat(),
        "profile": {"contact": "", "business": business.strip(),
                    "headcount": headcount, "areas_of_law": [], "notes": notes.strip()},
        "documents": [],
    }
    clients.append(client)
    save_clients(clients)
    return client
